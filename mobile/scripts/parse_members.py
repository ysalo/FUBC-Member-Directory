#!/usr/bin/env python3
"""Extract member records from the Ukrainian XLSX directory."""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel
except ImportError:
    print("This script requires openpyxl. Install it with: python -m pip install openpyxl", file=sys.stderr)
    raise SystemExit(2)


YEAR_PIVOT = 30  # 00-29 => 2000-2029; 30-99 => 1930-1999.
PHONE_RE = re.compile(r"(?<!\w)(?:\+?\d[\d\s().\-/]{5,}\d)(?!\w)")
DATE_RE = re.compile(r"^\s*(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})\s*$")
SUMMARY_RE = re.compile(r"всього|пенсіонер|сімей|одиноких|до\s+\d+", re.IGNORECASE)
LETTER_DIVIDER_RE = re.compile(r"^\s*~{2,}.*~{2,}\s*$")


def text_value(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_name(value: Any) -> dict[str, str | None] | None:
    """Parse surname, given name, and optional patronymic without combining fields."""
    parts = text_value(value).split()
    if len(parts) not in (2, 3):
        return None
    return {
        "first_name": normalize_apostrophes(parts[1]),
        "last_name": normalize_apostrophes(parts[0]),
        "patronymic": normalize_apostrophes(parts[2]) if len(parts) == 3 else None,
    }


def normalize_apostrophes(value: str) -> str:
    """Normalize straight, curly, modifier, and mis-entered quote marks."""
    return re.sub(r"['’‘ʼ`\"“”]", "’", value)


def categorize_issues(issues: list[dict[str, Any]]) -> tuple[dict[str, list[dict[str, Any]]], dict[str, int]]:
    """Group audit issues under category keys and return category totals."""
    grouped: dict[str, list[dict[str, Any]]] = {}
    for item in issues:
        issue = str(item.get("issue", "")).casefold()
        field = str(item.get("field", "")).casefold()
        if issue.startswith("interpreted "):
            category = "date_interpretation"
        elif "missing" in issue:
            category = "missing_data"
        elif field in {"birth_date", "membership_joined_at"} or any(token in issue for token in ("date", "numeric date", "invalid date")):
            category = "date_format"
        elif field in {"phone", "phone_number"} or "phone" in issue:
            category = "phone"
        elif field == "address" or "address" in issue:
            category = "address_quality"
        elif "name" in issue:
            category = "name_parse"
        else:
            category = "other"
        grouped.setdefault(category, []).append(item)
    grouped = dict(sorted(grouped.items()))
    return grouped, {category: len(items) for category, items in grouped.items()}


def parse_date(value: Any, *, epoch: datetime, number_format: str = "") -> tuple[str | None, str | None]:
    if value is None or text_value(value) == "":
        return None, None
    if isinstance(value, datetime):
        return value.date().isoformat(), None
    if isinstance(value, date):
        return value.isoformat(), None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = from_excel(value, epoch)
            if isinstance(parsed, datetime):
                parsed = parsed.date()
            if isinstance(parsed, date):
                return parsed.isoformat(), None
        except (OverflowError, TypeError, ValueError):
            pass
        return None, f"Could not interpret numeric date {value!r} (format {number_format!r})"

    raw = text_value(value)
    match = DATE_RE.fullmatch(raw)
    if not match:
        return None, f"Unrecognized date {raw!r}"
    first, second, year_text = (int(item) for item in match.groups())
    if len(match.group(3)) == 2:
        year = 2000 + year_text if year_text < YEAR_PIVOT else 1900 + year_text
    else:
        year = year_text

    # The workbook mostly uses MM-DD-YY. If the first number cannot be a
    # month, interpret that cell as DD-MM-YY and record the exception.
    if first > 12 and second <= 12:
        month, day = second, first
        note = f"Interpreted {raw!r} as DD-MM-YY"
    elif second > 12 and first <= 12:
        month, day = first, second
        note = None
    elif first <= 12 and second <= 12:
        month, day = first, second
        note = None
    else:
        return None, f"Invalid date {raw!r}"

    try:
        return date(year, month, day).isoformat(), note
    except ValueError:
        return None, f"Invalid date {raw!r}"


def phone_value(value: Any) -> str | None:
    raw = text_value(value)
    if not raw:
        return None
    matches = [match.group(0).strip() for match in PHONE_RE.finditer(raw)]
    matches = [candidate for candidate in matches if sum(char.isdigit() for char in candidate) >= 7]
    if matches:
        return max(matches, key=lambda candidate: sum(char.isdigit() for char in candidate))
    if sum(char.isdigit() for char in raw) >= 7:
        return raw
    return None


def find_member_sheet(workbook: Any) -> Any:
    for sheet in workbook.worksheets:
        first = text_value(sheet.cell(1, 1).value).casefold()
        second = text_value(sheet.cell(1, 2).value).casefold()
        if first in {"№", "номер"} and "член" in second:
            return sheet
    raise ValueError("Could not find the member sheet (expected headers № and Список Членів in row 1).")


def extract(input_path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(input_path, data_only=True)
    sheet = find_member_sheet(workbook)
    members: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    phone_audit: list[dict[str, Any]] = []

    for row_number in range(2, sheet.max_row + 1):
        name_cell = sheet.cell(row_number, 2)
        raw_name = text_value(name_cell.value)
        if not raw_name or LETTER_DIVIDER_RE.fullmatch(raw_name) or SUMMARY_RE.search(raw_name):
            continue

        parsed_name = parse_name(raw_name)
        if parsed_name is None:
            issues.append({"row": row_number, "member_name": raw_name, "issue": "Could not parse name in column B"})
            continue

        # The app requires gender for new people, but the spreadsheet has no
        # gender column. Keep it explicit and null so it can be completed later.
        member: dict[str, Any] = {**parsed_name, "gender": None, "email": None}
        address = text_value(sheet.cell(row_number, 3).value)
        if address:
            member["address"] = address
            if not any(character.isdigit() for character in address):
                issues.append({
                    "row": row_number,
                    "member_name": raw_name,
                    "field": "address",
                    "issue": "Address has no numeric street or unit detail; it may be city-only or a note",
                    "address_value": address,
                })

        home_phone = phone_value(sheet.cell(row_number, 4).value)
        phone_number = phone_value(sheet.cell(row_number, 5).value)
        if phone_number or home_phone:
            member["phone"] = phone_number or home_phone
        if home_phone:
            phone_audit.append({
                "row": row_number,
                "member_name": raw_name,
                "home_phone_found": True,
                "phone_number_source": "cellphone" if phone_number else "home_phone",
            })
        if not phone_number and not home_phone:
            issues.append({
                "row": row_number,
                "member_name": raw_name,
                "field": "phone_number",
                "issue": "No extractable phone number in either phone column",
            })

        for column, field in ((6, "birth_date"), (7, "membership_joined_at")):
            cell = sheet.cell(row_number, column)
            parsed_date, note = parse_date(cell.value, epoch=workbook.epoch, number_format=cell.number_format)
            if parsed_date:
                member[field] = parsed_date
            elif cell.value is not None and text_value(cell.value):
                issues.append({"row": row_number, "member_name": raw_name, "field": field, "issue": note or "Could not parse date"})
            elif cell.value is None or not text_value(cell.value):
                issues.append({"row": row_number, "member_name": raw_name, "field": field, "issue": "Missing date"})
            if note and parsed_date:
                issues.append({"row": row_number, "member_name": raw_name, "field": field, "issue": note})

        members.append(member)

    categorized_issues, issue_category_counts = categorize_issues(issues)
    audit = {
        "input_file": input_path.name,
        "sheet": sheet.title,
        "records_extracted": len(members),
        "home_phone_audit": phone_audit,
        "issues": categorized_issues,
        "issue_category_counts": issue_category_counts,
    }
    return members, audit


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract Ukrainian member records from an XLSX file."
    )
    parser.add_argument("xlsx_file", type=Path, help="Path to the original .xlsx workbook")
    parser.add_argument("--output-dir", type=Path, help="Directory for members.json and members_audit.json (default: beside the workbook)")
    args = parser.parse_args()
    input_path = args.xlsx_file.expanduser().resolve()
    if not input_path.is_file():
        parser.error(f"Input file does not exist: {input_path}")
    if input_path.suffix.casefold() != ".xlsx":
        parser.error("Input file must have an .xlsx extension")

    try:
        members, audit = extract(input_path)
    except Exception as error:
        print(f"Could not parse workbook: {error}", file=sys.stderr)
        return 1

    output_dir = args.output_dir.expanduser().resolve() if args.output_dir else input_path.parent
    output_dir.mkdir(parents=True, exist_ok=True)
    members_path = output_dir / "members.json"
    audit_path = output_dir / "members_audit.json"
    members_path.write_text(json.dumps(members, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(members)} members to {members_path}")
    print(f"Wrote audit report to {audit_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
