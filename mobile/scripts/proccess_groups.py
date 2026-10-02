#!/usr/bin/env python3
"""Extract member lists from the church's split (breakup) workbooks.

Each source workbook gets a folder under ``processed`` containing members.json,
members.csv, members_audit.json, and members_audit.csv. A combined
group_audit.json compares all groups with the main directory workbook.

Requires openpyxl. Legacy .xls inputs also require LibreOffice (so they can be
converted to .xlsx before reading).
"""

from __future__ import annotations

import argparse
import csv
from difflib import SequenceMatcher
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel
except ImportError:
    print("This script requires openpyxl. Install it with: python -m pip install openpyxl", file=sys.stderr)
    raise SystemExit(2)

# Reuse the established field/date/phone parsing rules for the main sheet.
SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))
import parse_members  # noqa: E402


MEMBER_COLUMNS = ["name", "patronymic", "gender", "email", "phone", "membership_joined_at", "birth_date", "address"]
AUDIT_COLUMNS = ["record_type", "row", "member_name", "field", "issue", "address_value", "home_phone_found", "phone_number_source"]
SKIP_NAME_RE = re.compile(r"всього|пенсіонер|сімей|одиноких|до\s+\d+", re.IGNORECASE)


def normalized_name(member: dict[str, Any]) -> str:
    return re.sub(r"\s+", " ", f"{member.get('name', '')} {member.get('patronymic', '')}").strip().casefold()


def member_key(member: dict[str, Any]) -> str:
    """Use date of birth to distinguish people who share a name."""
    name = normalized_name(member)
    birth_date = str(member.get("birth_date") or "").strip()
    return f"{name}\x1f{birth_date}" if birth_date else name


def name_similarity(left: str, right: str) -> float:
    """Compare normalized full names, also allowing name-token order changes."""
    compact_left = re.sub(r"[^\w]", "", left, flags=re.UNICODE)
    compact_right = re.sub(r"[^\w]", "", right, flags=re.UNICODE)
    direct = SequenceMatcher(None, compact_left, compact_right).ratio()
    left_tokens = " ".join(sorted(left.split()))
    right_tokens = " ".join(sorted(right.split()))
    reordered = SequenceMatcher(None, left_tokens, right_tokens).ratio()
    return max(direct, reordered)


def compare_member_fields(group_member: dict[str, Any], main_member: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """Compare the four requested identifying/contact fields when both exist."""
    evidence: dict[str, dict[str, Any]] = {}
    for field in ("address", "phone", "birth_date", "membership_joined_at"):
        left = str(group_member.get(field) or "").strip().casefold()
        right = str(main_member.get(field) or "").strip().casefold()
        if field == "phone":
            left = re.sub(r"\D", "", left)
            right = re.sub(r"\D", "", right)
        elif field == "address":
            left = re.sub(r"[^\w]", "", left, flags=re.UNICODE)
            right = re.sub(r"[^\w]", "", right, flags=re.UNICODE)
        if not left or not right:
            evidence[field] = {"similarity": None, "similar": False, "compared": False}
            continue
        if field in {"birth_date", "membership_joined_at"}:
            score = 1.0 if left == right else 0.0
        else:
            score = SequenceMatcher(None, left, right).ratio()
        evidence[field] = {"similarity": round(score, 4), "similar": score >= 0.90, "compared": True}
    return evidence


def as_xlsx(source: Path, tempdir: Path) -> Path:
    if source.suffix.casefold() == ".xlsx":
        return source
    result = subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tempdir), str(source)],
        capture_output=True, text=True,
    )
    converted = tempdir / f"{source.stem}.xlsx"
    if result.returncode != 0 or not converted.exists():
        detail = result.stderr.strip() or result.stdout.strip() or "conversion failed"
        raise RuntimeError(f"Could not convert legacy XLS file {source.name}: {detail}")
    return converted


def extract_group(source: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    with tempfile.TemporaryDirectory(prefix="breakup-xlsx-") as temporary:
        workbook_path = as_xlsx(source, Path(temporary))
        workbook = load_workbook(workbook_path, data_only=True)
        sheet = parse_members.find_member_sheet(workbook)
        members: list[dict[str, Any]] = []
        issues: list[dict[str, Any]] = []

        for row_number in range(2, sheet.max_row + 1):
            raw_name = parse_members.text_value(sheet.cell(row_number, 2).value)
            if not raw_name:
                continue
            if parse_members.LETTER_DIVIDER_RE.fullmatch(raw_name):
                continue
            if SKIP_NAME_RE.search(raw_name):
                continue
            parsed = parse_members.parse_name(raw_name)
            if parsed is None:
                issues.append({"row": row_number, "member_name": raw_name, "issue": "Could not parse name in column B"})
                continue

            member: dict[str, Any] = {**parsed, "gender": None, "email": None}
            address = parse_members.text_value(sheet.cell(row_number, 3).value)
            if address:
                member["address"] = address
            phone = parse_members.phone_value(sheet.cell(row_number, 4).value)
            if phone:
                member["phone"] = phone
            else:
                issues.append({"row": row_number, "member_name": raw_name, "field": "phone", "issue": "No extractable phone number"})

            for column, field in ((5, "birth_date"), (6, "membership_joined_at")):
                cell = sheet.cell(row_number, column)
                parsed_date, note = parse_members.parse_date(
                    cell.value, epoch=workbook.epoch, number_format=cell.number_format
                )
                if parsed_date:
                    member[field] = parsed_date
                elif cell.value is not None and parse_members.text_value(cell.value):
                    issues.append({"row": row_number, "member_name": raw_name, "field": field, "issue": note or "Could not parse date"})
                else:
                    issues.append({"row": row_number, "member_name": raw_name, "field": field, "issue": "Missing date"})
                if note and parsed_date:
                    issues.append({"row": row_number, "member_name": raw_name, "field": field, "issue": note})
            members.append(member)

        categorized_issues, issue_category_counts = parse_members.categorize_issues(issues)
        return members, {
            "input_file": source.name,
            "sheet": sheet.title,
            "records_extracted": len(members),
            "issues": categorized_issues,
            "issue_category_counts": issue_category_counts,
        }


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows({key: "" if row.get(key) is None else row.get(key, "") for key in columns} for row in rows)


def write_group_outputs(folder: Path, members: list[dict[str, Any]], audit: dict[str, Any]) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "members.json").write_text(json.dumps(members, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    member_columns = MEMBER_COLUMNS + sorted({key for row in members for key in row if key not in MEMBER_COLUMNS})
    write_csv(folder / "members.csv", members, member_columns)
    (folder / "members_audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    audit_rows = [
        {"record_type": "issue", **issue}
        for category_issues in audit["issues"].values()
        for issue in category_issues
    ]
    write_csv(folder / "members_audit.csv", audit_rows, AUDIT_COLUMNS)


def main() -> int:
    default_original = Path("original")
    default_source = default_original / "groups"
    default_main = default_original / "main.xlsx"
    default_output = Path("processed")
    parser = argparse.ArgumentParser(
        description="Extract each breakup workbook to its own output folder and compare group members with the main directory.",
        epilog=("Example: python mobile/scripts/proccess_groups.py "
                "--source-dir original/groups "
                "--main-workbook original/main.xlsx "
                "--output-dir processed"),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--source-dir", type=Path, default=default_source, help="Folder containing the group .xls/.xlsx files")
    parser.add_argument("--main-workbook", type=Path, default=default_main, help="Main directory workbook used for comparison")
    parser.add_argument("--output-dir", type=Path, default=default_output, help="Processed output root (group files go in its groups/ subfolder)")
    parser.add_argument(
        "--fuzzy-threshold", type=float, default=0.90,
        help="Minimum name similarity from 0 to 1 for fuzzy matching (default: 0.90)",
    )
    args = parser.parse_args()
    source_dir = args.source_dir.expanduser().resolve()
    main_path = args.main_workbook.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    group_output_dir = output_dir / "groups"
    if not source_dir.is_dir():
        parser.error(f"Source directory does not exist: {source_dir}")
    if not main_path.is_file():
        parser.error(f"Main workbook does not exist: {main_path}")
    if not 0 <= args.fuzzy_threshold <= 1:
        parser.error("--fuzzy-threshold must be between 0 and 1")

    sources = sorted(path for path in source_dir.iterdir() if path.is_file() and path.suffix.casefold() in {".xls", ".xlsx"})
    sources = [path for path in sources if not path.name.startswith("~$") and output_dir not in path.parents]
    if not sources:
        parser.error(f"No .xls or .xlsx group workbooks found in {source_dir}")

    try:
        main_members, main_audit = parse_members.extract(main_path)
        group_entries: list[dict[str, Any]] = []
        all_group_members: list[dict[str, Any]] = []
        group_names_by_member: dict[str, list[str]] = {}
        group_issue_category_totals: dict[str, int] = {}
        for source in sources:
            members, audit = extract_group(source)
            group_name = source.stem.strip()
            folder = group_output_dir / group_name
            write_group_outputs(folder, members, audit)
            issue_count = sum(len(items) for items in audit["issues"].values())
            group_entries.append({"group": group_name, "input_file": source.name, "records_extracted": len(members), "issues_count": issue_count, "issue_category_counts": audit["issue_category_counts"], "output_folder": str(folder)})
            for category, count in audit["issue_category_counts"].items():
                group_issue_category_totals[category] = group_issue_category_totals.get(category, 0) + count
            all_group_members.extend(members)
            for member in members:
                key = member_key(member)
                group_names_by_member.setdefault(key, []).append(group_name)

        main_by_name: dict[str, list[dict[str, Any]]] = {}
        groups_by_name: dict[str, list[dict[str, Any]]] = {}
        groups_by_identity: dict[str, list[dict[str, Any]]] = {}
        for member in main_members:
            main_by_name.setdefault(normalized_name(member), []).append(member)
        for member in all_group_members:
            groups_by_name.setdefault(normalized_name(member), []).append(member)
            groups_by_identity.setdefault(member_key(member), []).append(member)
        exact_names = set(main_by_name) & set(groups_by_name)
        unmatched_main_names = set(main_by_name) - exact_names
        unmatched_group_names = set(groups_by_name) - exact_names
        fuzzy_matches: list[dict[str, Any]] = []
        ambiguous_matches: list[dict[str, Any]] = []
        fuzzy_matched_main_names: set[str] = set()
        for group_name in sorted(unmatched_group_names):
            ranked = []
            for main_name in unmatched_main_names:
                group_member = groups_by_name[group_name][0]
                main_member = main_by_name[main_name][0]
                name_score = name_similarity(normalized_name(group_member), normalized_name(main_member))
                field_evidence = compare_member_fields(group_member, main_member)
                field_match_count = sum(item["similar"] for item in field_evidence.values())
                same_given_and_surname = (
                    re.sub(r"\s+", " ", str(group_member.get("name", ""))).strip().casefold()
                    == re.sub(r"\s+", " ", str(main_member.get("name", ""))).strip().casefold()
                )
                same_dates = all(
                    group_member.get(field) and main_member.get(field)
                    and str(group_member[field]) == str(main_member[field])
                    for field in ("birth_date", "membership_joined_at")
                )
                exact_name_and_dates = same_given_and_surname and same_dates
                name_and_fields = name_score > args.fuzzy_threshold and field_match_count >= 2
                if exact_name_and_dates or name_and_fields:
                    rule = (
                        "given_name_surname_and_exact_birth_and_membership_dates"
                        if exact_name_and_dates else "name_similarity_and_two_similar_fields"
                    )
                    ranked.append((field_match_count, name_score, main_name, field_evidence, rule))
            ranked.sort(key=lambda item: (item[0], item[1]), reverse=True)
            if not ranked:
                continue
            second = ranked[1] if len(ranked) > 1 else None
            if second and ranked[0][0] == second[0] and ranked[0][1] - second[1] < 0.03:
                ambiguous_matches.append({
                    "group_member": groups_by_name[group_name][0],
                    "candidates": [
                        {"main_member": main_by_name[candidate][0], "name_similarity": round(score, 4), "similar_fields": count, "field_similarities": evidence}
                        for count, score, candidate, evidence, _rule in ranked[:3]
                    ],
                })
                continue
            field_match_count, score, main_name, field_evidence, match_rule = ranked[0]
            fuzzy_matched_main_names.add(main_name)
            fuzzy_matches.append({
                "status": "matched",
                "group_member": groups_by_name[group_name][0],
                "main_member": main_by_name[main_name][0],
                "name_similarity": round(score, 4),
                "similar_fields": field_match_count,
                "field_similarities": field_evidence,
                "match_rule": match_rule,
            })

        matched_main_names = exact_names | fuzzy_matched_main_names
        fuzzy_group_names = {normalized_name(item["group_member"]) for item in fuzzy_matches}
        missing_from_groups = [rows[0] for name, rows in main_by_name.items() if name not in matched_main_names]
        absent_from_main = [rows[0] for name, rows in groups_by_name.items() if name not in exact_names and name not in fuzzy_group_names]

        # Make likely spelling/patronymic differences easy to review. High
        # confidence fuzzy matches are highlighted, and unresolved entries get
        # candidate suggestions when the first+last name matches or the full
        # name is still reasonably similar. Suggestions do not alter counts.
        likely_name_match_issues: list[dict[str, Any]] = [
            {
                "status": "matched",
                "group_member": item["group_member"],
                "main_member_candidate": item["main_member"],
                "name_similarity": item["name_similarity"],
                "similar_fields": item["similar_fields"],
                "field_similarities": item["field_similarities"],
                "review_reason": f"Matched by rule: {item['match_rule']}; verify any name spelling difference.",
            }
            for item in fuzzy_matches
        ]
        candidate_name_matches_for_review: list[dict[str, Any]] = []
        for group_member in absent_from_main:
            group_display_name = re.sub(r"\s+", " ", str(group_member.get("name", ""))).strip().casefold()
            candidates: list[tuple[float, dict[str, Any]]] = []
            for main_member in missing_from_groups:
                main_display_name = re.sub(r"\s+", " ", str(main_member.get("name", ""))).strip().casefold()
                score = name_similarity(normalized_name(group_member), normalized_name(main_member))
                field_evidence = compare_member_fields(group_member, main_member)
                if group_display_name == main_display_name or score >= 0.75:
                    candidates.append((score, main_member))
            candidates.sort(key=lambda pair: pair[0], reverse=True)
            for score, candidate in candidates[:3]:
                candidate_name_matches_for_review.append({
                    "status": "candidate_for_review",
                    "group_member": group_member,
                    "main_member_candidate": candidate,
                    "name_similarity": round(score, 4),
                    "similar_fields": sum(item["similar"] for item in compare_member_fields(group_member, candidate).values()),
                    "field_similarities": compare_member_fields(group_member, candidate),
                    "review_reason": (
                        "Same given-name/surname but patronymic is missing or differs."
                        if re.sub(r"\s+", " ", str(group_member.get("name", ""))).strip().casefold()
                        == re.sub(r"\s+", " ", str(candidate.get("name", ""))).strip().casefold()
                        else "Names are similar but did not meet the automatic fuzzy-match rule."
                    ),
                })

        potential_group_matches_for_main_unmatched: list[dict[str, Any]] = []
        for main_member in missing_from_groups:
            potential: list[tuple[float, int, dict[str, Any]]] = []
            seen_group_names: set[str] = set()
            for group_member in all_group_members:
                group_key = member_key(group_member)
                if group_key in seen_group_names:
                    continue
                seen_group_names.add(group_key)
                name_score = name_similarity(normalized_name(main_member), normalized_name(group_member))
                if main_member.get("birth_date") and group_member.get("birth_date") and main_member["birth_date"] != group_member["birth_date"]:
                    continue
                evidence = compare_member_fields(main_member, group_member)
                exact_field_count = sum(item["similarity"] == 1.0 for item in evidence.values() if item["compared"])
                phone_match = evidence["phone"]["similarity"] == 1.0 and evidence["phone"]["compared"]
                same_given_and_surname = (
                    re.sub(r"\s+", " ", str(main_member.get("name", ""))).strip().casefold()
                    == re.sub(r"\s+", " ", str(group_member.get("name", ""))).strip().casefold()
                )
                if phone_match or (same_given_and_surname and name_score >= 0.75) or (name_score >= 0.60 and exact_field_count >= 2):
                    potential.append((name_score, exact_field_count, {
                        "group_member": group_member,
                        "groups": sorted(set(group_names_by_member.get(group_key, []))),
                        "name_similarity": round(name_score, 4),
                        "exact_fields": exact_field_count,
                        "field_similarities": evidence,
                    }))
            potential.sort(key=lambda item: (item[0], item[1]), reverse=True)
            if potential:
                potential_group_matches_for_main_unmatched.append({
                    "main_member": main_member,
                    "candidates": [item[2] for item in potential[:5]],
                })
        duplicates_in_groups = [
            {
                "member": rows[0],
                "occurrences": len(rows),
                "groups": sorted(set(group_names_by_member.get(name, []))),
            }
            for name, rows in groups_by_identity.items() if len(rows) > 1
        ]

        report = {
            "main_workbook": str(main_path),
            "main_issue_category_counts": main_audit["issue_category_counts"],
            "group_issue_category_totals": dict(sorted(group_issue_category_totals.items())),
            "group_count": len(group_entries),
            "groups": group_entries,
            "group_member_total": len(all_group_members),
            "unique_group_member_total": len(groups_by_identity),
            "main_member_total": len(main_members),
            "members_in_main_not_in_any_group": missing_from_groups,
            "members_in_groups_not_in_main": absent_from_main,
            "fuzzy_name_matches": fuzzy_matches,
            "ambiguous_name_matches_for_review": ambiguous_matches,
            "likely_name_match_issues": likely_name_match_issues,
            "candidate_name_matches_for_review": candidate_name_matches_for_review,
            "potential_group_matches_for_main_unmatched": potential_group_matches_for_main_unmatched,
            "potential_group_match_search": {
                "criteria": "exact phone, or matching given name/surname with name similarity >= 0.75, or name similarity >= 0.60 plus two exact fields",
                "date_fields_require_exact_equality": True,
            },
            "fuzzy_matching": {
                "method": "SequenceMatcher on normalized full names and sorted name tokens",
                "name_similarity_strictly_greater_than": args.fuzzy_threshold,
                "required_similar_fields": 2,
                "field_threshold": 0.90,
                "fields": ["address", "phone", "birth_date", "membership_joined_at"],
                "additional_match_rule": "given name and surname match, and both birth_date and membership_joined_at are present and exactly equal",
                "ambiguity_margin": 0.03,
            },
            "duplicate_members_across_groups": duplicates_in_groups,
            "comparison_counts": {
                "main_not_in_groups": len(missing_from_groups),
                "groups_not_in_main": len(absent_from_main),
                "fuzzy_matches": len(fuzzy_matches),
                "ambiguous_matches_for_review": len(ambiguous_matches),
                "likely_name_match_issues": len(likely_name_match_issues),
                "candidate_name_matches_for_review": len(candidate_name_matches_for_review),
                "main_unmatched_with_potential_group_candidates": len(potential_group_matches_for_main_unmatched),
                "duplicate_group_occurrences": sum(item["occurrences"] - 1 for item in duplicates_in_groups),
            },
        }
        output_dir.mkdir(parents=True, exist_ok=True)
        report_path = output_dir / "group_audit.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except Exception as error:
        print(f"Could not process breakup workbooks: {error}", file=sys.stderr)
        return 1

    for entry in group_entries:
        print(f"{entry['records_extracted']:>4} members: {entry['input_file']} -> {entry['output_folder']}")
    print(f"Wrote combined audit: {report_path}")
    print(f"Group members: {len(all_group_members)}; main members: {len(main_members)}; unmatched main/group: {len(missing_from_groups)}/{len(absent_from_main)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
